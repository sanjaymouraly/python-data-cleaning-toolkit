"""Exporter tests: round-trips, overwrite safety, formula guard, atomicity."""

from __future__ import annotations

import os

import numpy as np
import pandas as pd
import pytest
from hypothesis import given, settings
from hypothesis import strategies as st

from data_cleaning_toolkit import SafetyError
from data_cleaning_toolkit.exporters.exporters import (
    export_csv,
    export_dataframe,
    export_json,
    export_xlsx,
    sanitize_filename,
)


@pytest.fixture()
def typed_df():
    return pd.DataFrame(
        {
            "text": ["alpha", "beta", None],
            "num": [1.5, 2.0, np.nan],
            "int": [10, 20, 30],
            "when": pd.to_datetime(["2026-01-01", "2026-06-15", "2026-12-31"]),
        }
    )


class TestCsvExport:
    def test_roundtrip(self, typed_df, tmp_path):
        path = str(tmp_path / "out.csv")
        res = export_csv(typed_df, path)
        assert res.validated
        back = pd.read_csv(path, dtype=object, keep_default_na=False)
        assert len(back) == 3
        assert list(back.columns) == list(typed_df.columns)

    def test_overwrite_refused(self, typed_df, tmp_path):
        path = str(tmp_path / "out.csv")
        export_csv(typed_df, path)
        with pytest.raises(SafetyError):
            export_csv(typed_df, path)

    def test_overwrite_allowed_explicitly(self, typed_df, tmp_path):
        path = str(tmp_path / "out.csv")
        export_csv(typed_df, path)
        export_csv(typed_df, path, allow_overwrite=True)  # no raise

    def test_formula_injection_guard(self, tmp_path):
        df = pd.DataFrame({"cmd": ["=SUM(A1:A2)", "@cmd", "+calc", "normal"]})
        path = str(tmp_path / "guard.csv")
        res = export_csv(df, path)
        assert res.guarded_cells == 3
        raw = open(path, encoding="utf-8").read()
        assert "'=SUM(A1:A2)" in raw
        assert "'@cmd" in raw

    def test_guard_can_be_disabled(self, tmp_path):
        df = pd.DataFrame({"cmd": ["=SUM(A1)"]})
        path = str(tmp_path / "noguard.csv")
        res = export_csv(df, path, formula_guard=False)
        assert res.guarded_cells == 0

    def test_atomic_no_partial_file_on_failure(self, typed_df, tmp_path, monkeypatch):
        """If writing fails midway, no partial file (or temp file) may remain."""
        path = str(tmp_path / "boom.csv")

        def broken_to_csv(self, *args, **kwargs):
            raise RuntimeError("disk exploded mid-write")

        monkeypatch.setattr(pd.DataFrame, "to_csv", broken_to_csv)
        with pytest.raises(RuntimeError):
            export_csv(typed_df, path, validate=False)
        leftovers = [f for f in os.listdir(tmp_path) if f.startswith(".dtk_")]
        assert leftovers == []
        assert not os.path.exists(path)
    def test_leading_zeros_preserved(self, tmp_path):
        df = pd.DataFrame({"zip": ["00123"]})
        path = str(tmp_path / "z.csv")
        export_csv(df, path)
        assert "00123" in open(path, encoding="utf-8").read()

    def test_directory_created(self, typed_df, tmp_path):
        target = tmp_path / "deep" / "nested" / "out.csv"
        export_csv(typed_df, str(target))
        assert target.exists()


class TestXlsxExport:
    def test_roundtrip(self, typed_df, tmp_path):
        path = str(tmp_path / "out.xlsx")
        res = export_xlsx(typed_df, path)
        assert res.validated
        from openpyxl import load_workbook

        wb = load_workbook(path)
        ws = wb.active
        assert ws.max_row == 4  # header + 3
        wb.close()

    def test_overwrite_refused(self, typed_df, tmp_path):
        path = str(tmp_path / "o.xlsx")
        export_xlsx(typed_df, path)
        with pytest.raises(SafetyError):
            export_xlsx(typed_df, path)

    def test_dangerous_text_guarded(self, tmp_path):
        df = pd.DataFrame({"t": ["=HYPERLINK(\"http://evil\")"]})
        path = str(tmp_path / "g.xlsx")
        res = export_xlsx(df, path)
        assert res.guarded_cells == 1

    def test_na_becomes_empty_cell(self, tmp_path):
        df = pd.DataFrame({"v": [1.0, np.nan]})
        path = str(tmp_path / "na.xlsx")
        export_xlsx(df, path)
        from openpyxl import load_workbook

        wb = load_workbook(path, data_only=True)
        ws = wb.active
        assert ws.cell(row=3, column=1).value in (None, "")
        wb.close()


class TestJsonExport:
    def test_roundtrip(self, typed_df, tmp_path):
        path = str(tmp_path / "o.json")
        res = export_json(typed_df, path)
        assert res.validated
        import json

        records = json.load(open(path, encoding="utf-8"))
        assert len(records) == 3

    def test_overwrite_refused(self, typed_df, tmp_path):
        path = str(tmp_path / "o.json")
        export_json(typed_df, path)
        with pytest.raises(SafetyError):
            export_json(typed_df, path)


class TestHelpers:
    def test_sanitize_filename(self):
        assert sanitize_filename("../../etc/passwd") == "passwd"
        assert sanitize_filename("my data.csv") == "my data.csv"
        assert sanitize_filename("a/b\\c") == "a b c" or "/" not in sanitize_filename("a/b\\c")
        assert sanitize_filename("") == "output"

    def test_export_dataframe_dispatch(self, typed_df, tmp_path, simple_cfg):
        cfg = simple_cfg([])
        cfg = cfg.model_copy(
            update={
                "export": cfg.export.model_copy(
                    update={"fmt": "xlsx", "directory": str(tmp_path)}
                )
            }
        )
        res = export_dataframe(typed_df, str(tmp_path), "data", cfg)
        assert res.fmt == "xlsx" and res.validated

    def test_export_dataframe_none(self, typed_df, tmp_path, simple_cfg):
        cfg = simple_cfg([])
        cfg = cfg.model_copy(update={"export": cfg.export.model_copy(update={"fmt": "none"})})
        assert export_dataframe(typed_df, str(tmp_path), "d", cfg) is None


class TestPropertyBased:
    @given(st.text(alphabet="ab=+@ \t", max_size=12))
    def test_guard_never_corrupts_safe_text(self, s):
        from data_cleaning_toolkit.exporters.exporters import guard_cell

        out, guarded = guard_cell(s, True)
        if s.startswith(("=", "+", "@", "\t", "\r")):
            assert guarded and out == "'" + s
        else:
            assert out == s and not guarded

    @given(st.integers(min_value=-10**9, max_value=10**9))
    @settings(deadline=None)  # file I/O must not flake under full-suite load
    def test_csv_numeric_roundtrip(self, n):
        import tempfile

        df = pd.DataFrame({"v": [n]})
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "dtk_prop.csv")
            export_csv(df, path, allow_overwrite=True, validate=False)
            back = pd.read_csv(path)
        assert int(back["v"].iloc[0]) == n
